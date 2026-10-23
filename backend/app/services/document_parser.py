import csv
import json
from dataclasses import dataclass
from pathlib import Path

from docx import Document as WordDocument
from fastapi import HTTPException
from pypdf import PdfReader

from app.services.parsers.cdr_parser import is_cdr_csv, parse_cdr_file
from app.services.parsers.financial_parser import is_financial_csv, parse_financial_file

PARSER_VERSION = "tattva-text-v1"


@dataclass(frozen=True)
class TextUnit:
    unit_type: str
    unit_number: int
    text: str


@dataclass(frozen=True)
class ParsedText:
    text: str
    units: list[dict[str, int | str]]


def _normalize(text: str) -> str:
    # Only normalize line endings. Offsets in chunks continue to address this
    # stored extracted-text version exactly.
    return text.replace("\r\n", "\n").replace("\r", "\n")


def parse_document(path: Path, extension: str, max_chars: int) -> ParsedText:
    try:
        if extension == ".pdf":
            parsed_units = _parse_pdf(path)
        elif extension == ".txt":
            parsed_units = [TextUnit("record", 1, path.read_text(encoding="utf-8-sig"))]
        elif extension == ".docx":
            parsed_units = _parse_docx(path)
        elif extension == ".json":
            value = json.loads(path.read_text(encoding="utf-8-sig"))
            parsed_units = [TextUnit("record", 1, json.dumps(value, ensure_ascii=False, indent=2))]
        elif extension == ".csv":
            parsed_units = _parse_csv(path)
        else:
            raise HTTPException(status_code=415, detail="Unsupported file type.")
    except HTTPException:
        raise
    except (UnicodeError, json.JSONDecodeError, csv.Error, OSError, ValueError) as exc:
        raise HTTPException(status_code=422, detail="The uploaded file could not be parsed.") from exc
    except Exception as exc:
        # Parser exceptions are not returned to clients because they may expose
        # paths, parser internals or source data.
        raise HTTPException(status_code=422, detail="The uploaded file could not be parsed.") from exc

    text_parts: list[str] = []
    units: list[dict[str, int | str]] = []
    cursor = 0
    for unit in parsed_units:
        normalized = _normalize(unit.text)
        if not normalized:
            continue
        if text_parts:
            text_parts.append("\n")
            cursor += 1
        start = cursor
        text_parts.append(normalized)
        cursor += len(normalized)
        units.append(
            {
                "unit_type": unit.unit_type,
                "unit_number": unit.unit_number,
                "start_offset": start,
                "end_offset": cursor,
            }
        )
    text = "".join(text_parts)
    if not text.strip():
        raise HTTPException(status_code=422, detail="No extractable text was found in the uploaded file.")
    if len(text) > max_chars:
        raise HTTPException(status_code=413, detail="Extracted text exceeds the configured processing limit.")
    return ParsedText(text=text, units=units)


def _parse_pdf(path: Path) -> list[TextUnit]:
    reader = PdfReader(path, strict=False)
    if reader.is_encrypted:
        raise ValueError("Encrypted PDFs are unsupported.")
    return [TextUnit("page", index + 1, page.extract_text() or "") for index, page in enumerate(reader.pages)]


def _parse_docx(path: Path) -> list[TextUnit]:
    document = WordDocument(path)
    content: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text:
            content.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            content.append("\t".join(cell.text.replace("\n", " ") for cell in row.cells))
    return [TextUnit("record", 1, "\n".join(content))]


def _parse_csv(path: Path) -> list[TextUnit]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        if not reader.fieldnames:
            raise ValueError("CSV must have a header row.")
        fieldnames = list(reader.fieldnames)

    # 1. Check for Telecom CDR format
    if is_cdr_csv(fieldnames):
        full_text, records, summary = parse_cdr_file(path)
        units = [TextUnit("record", idx, line) for idx, line in enumerate(records, start=1)]
        return units

    # 2. Check for Banking / Hawala financial ledger format
    if is_financial_csv(fieldnames):
        full_text, records, summary = parse_financial_file(path)
        units = [TextUnit("record", idx, line) for idx, line in enumerate(records, start=1)]
        return units

    # 3. Standard CSV generic ingestion
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        units = []
        for row_number, row in enumerate(reader, start=1):
            clean = {key: value for key, value in row.items() if key is not None}
            units.append(TextUnit("record", row_number, json.dumps(clean, ensure_ascii=False, sort_keys=True)))
    return units


def split_chunks(text: str, units: list[dict], max_chars: int, overlap_chars: int) -> list[dict]:
    if overlap_chars >= max_chars:
        raise ValueError("Chunk overlap must be smaller than the chunk size.")
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            floor = start + int(max_chars * 0.65)
            boundary = max(text.rfind("\n", floor, end), text.rfind(" ", floor, end), text.rfind("\t", floor, end))
            if boundary > floor:
                end = boundary
        if end <= start:
            end = min(start + max_chars, len(text))
        unit = next(
            (u for u in units if u["start_offset"] <= start < u["end_offset"]),
            None,
        )
        chunks.append(
            {
                "text": text[start:end],
                "start_offset": start,
                "end_offset": end,
                "unit_type": unit["unit_type"] if unit else "record",
                "unit_number": unit["unit_number"] if unit else 1,
            }
        )
        if end >= len(text):
            break
        start = max(start + 1, end - overlap_chars)
    return chunks
